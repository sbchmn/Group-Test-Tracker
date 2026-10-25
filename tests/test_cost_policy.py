import json
import html
import re
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from flask_migrate import downgrade, upgrade
from openpyxl import load_workbook
from sqlalchemy import inspect, text
from sqlalchemy.orm import Session
from sqlalchemy.orm.exc import StaleDataError

from app import create_app, db
from app import costs
from app.export import generate_test_export
from app.models import GroupTest, NotificationTemplate, Participation, PaymentSubmission, User
from app.routes import _telegram_status_summary_for_user


class CostPolicyTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = create_app({'TESTING': True, 'WTF_CSRF_ENABLED': False,
                              'SQLALCHEMY_DATABASE_URI': 'sqlite:///' + str(Path(self.tmp.name) / 'costs.db'),
                              'NOTIFICATION_LOG_PATH': str(Path(self.tmp.name) / 'notifications.log')})
        self.client = self.app.test_client()
        with self.app.app_context():
            db.create_all()
            users = [User(username=name, email=name + '@example.test', is_admin=name == 'admin')
                     for name in ['admin', 'donor', 'member', 'late', 'other']]
            for user in users:
                user.set_password('secret')
            db.session.add_all(users)
            db.session.flush()
            test = GroupTest(title='Frozen pricing test', created_by=users[0].id, status='recruiting',
                             total_lab_cost=600, refund_per_donor=20)
            db.session.add(test)
            db.session.flush()
            self.ids = {}
            self.users = {user.username: user.id for user in users}
            for user in users[1:]:
                part = Participation(group_test_id=test.id, user_id=user.id, name=user.username,
                                     vial_donor=user.username == 'donor', approved=False)
                db.session.add(part)
                db.session.flush()
                self.ids[user.username] = part.id
            self.test_id = test.id
            db.session.commit()
        self.login('admin')

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()
        self.tmp.cleanup()

    def login(self, name):
        self.client.get('/logout')
        self.client.post('/login', data={'username': name, 'password': 'secret'})

    def approve(self, name, queue=False, **data):
        path = '/admin/action-queue/approve/' if queue else '/admin/approve-request/'
        return self.client.post(path + str(self.ids[name]), data=data)

    def edit_test(self, status, **extra):
        payload = {'title': 'Frozen pricing test', 'status': status, 'total_lab_cost': '600',
                   'shipping_cost': '0', 'donor_shipping_cost': '0', 'refund_per_donor': '20',
                   'donor_shipping_reimbursement': 'credit', 'donor_shipping_reimbursed_by_id': '0'}
        payload.update(extra)
        with patch('app.routes._notify_status_change_after_commit'):
            return self.client.post(f'/admin/edit-test/{self.test_id}', data=payload)

    def charges(self):
        with self.app.app_context():
            return {p.name: p.amount_owed for p in Participation.query.filter_by(group_test_id=self.test_id, approved=True).all()}

    def freeze(self):
        self.approve('donor')
        self.approve('member')
        self.assertEqual(self.edit_test('ready_for_payment').status_code, 302)
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            self.assertTrue(test.costs_frozen)
            self.assertIsNotNone(test.costs_frozen_at)
            self.assertEqual(test.costs_frozen_by_id, self.users['admin'])

    def explanation(self, path):
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)
        match = re.search(r'<details\b([^>]*)>\s*<summary>How the split works</summary>(.*?)</details>', response.get_data(as_text=True), re.S)
        self.assertIsNotNone(match)
        self.assertNotRegex(match.group(1), r'\bopen\b')
        return ' '.join(html.unescape(re.sub(r'<[^>]+>', ' ', match.group(2))).split())

    def test_detailed_explanation_uses_frozen_inputs_and_shipping_credit_on_both_pages(self):
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            test.total_lab_cost = 100
            test.shipping_cost = 20
            test.donor_shipping_cost = 15
            test.donor_shipping_reimbursement = 'credit'
            db.session.commit()
        self.approve('donor')
        self.approve('member')
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            costs.freeze(test, self.users['admin'])
            test.status = 'ready_for_payment'
            db.session.commit()
            test.total_lab_cost = 9999
            test.donor_shipping_cost = 999
            test.refund_per_donor = 99
            test.donor_shipping_reimbursement = 'participant'
            db.session.commit()
        for path in (f'/test/{self.test_id}', f'/admin/manage-participants/{self.test_id}'):
            with self.subTest(path=path):
                text = self.explanation(path)
                self.assertIn('lab cost + shipping + donor shipping = $100.00 + $20.00 + $15.00 = $135.00', text)
                self.assertIn('$135.00 ÷ 2 cost-sharing participants = $67.50 each', text)
                self.assertIn('1 donor × $20.00 credit per donor = $20.00', text)
                self.assertIn('$20.00 ÷ 1 non-donor = $20.00', text)
                self.assertIn('$67.50 base share + $20.00 uplift = $87.50', text)
                self.assertIn('$67.50 base share − $20.00 donor credit − $15.00 donor-shipping credit = $32.50', text)
                self.assertIn('Reimbursement policy: credit to the donor.', text)
                self.assertNotIn('$9999.00', text)
                self.assertNotIn('$999.00', text)
        self.assertEqual(self.charges(), {'donor': 32.5, 'member': 87.5})

    def test_explanation_handles_no_participants_no_donors_and_all_donors(self):
        text = self.explanation(f'/test/{self.test_id}')
        self.assertIn('no approved cost-sharing participants yet', text)
        self.assertNotIn('÷ 0', text)
        with self.app.app_context():
            db.session.get(GroupTest, self.test_id).total_lab_cost = 883
            db.session.commit()
        self.approve('member')
        text = self.explanation(f'/test/{self.test_id}')
        self.assertIn('$883.00 ÷ 1 cost-sharing participant = $883.00 each', text)
        self.assertIn('no vial donors', text)
        self.assertNotIn('$20.00', text)
        self.assertNotIn('÷ 0', text)
        with self.app.app_context():
            db.session.get(Participation, self.ids['member']).approved = False
            db.session.commit()
        self.approve('donor')
        text = self.explanation(f'/test/{self.test_id}')
        self.assertIn('no non-donors to fund an uplift', text)
        self.assertIn('$883.00 base share − $20.00 donor credit = $863.00', text)
        self.assertNotIn('÷ 0', text)

    def test_manage_and_queue_approval_reprice_entire_recruiting_cohort_quietly(self):
        with patch('app.routes.send_notification_message') as notify:
            self.approve('donor')
            self.assertEqual(self.charges()['donor'], 580)
            self.approve('member', queue=True)
            self.assertEqual(self.charges(), {'donor': 280, 'member': 320})
            self.approve('other')
            self.assertEqual(self.charges(), {'donor': 180, 'member': 210, 'other': 210})
        notify.assert_not_called()

    def test_bulk_approval_recalculates_once_and_frozen_bulk_uses_default(self):
        self.client.post('/admin/action-queue/approve-selected', data={'part_ids': [self.ids['donor'], self.ids['member']]})
        self.assertEqual(self.charges(), {'donor': 280, 'member': 320})
        self.edit_test('ready_for_payment')
        self.client.post('/admin/action-queue/approve-selected', data={'part_ids': [self.ids['late'], self.ids['other']]})
        self.assertEqual(self.charges(), {'donor': 280, 'member': 320, 'late': 320, 'other': 320})

    def test_late_buy_in_defaults_and_override_never_reprice_existing_members(self):
        self.freeze()
        self.approve('late', queue=True)
        self.approve('other', buy_in_amount='75.25')
        self.assertEqual(self.charges(), {'donor': 280, 'member': 320, 'late': 320, 'other': 75.25})
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            data = test.calculate_costs()
            self.assertEqual((data['total_participants'], data['total_non_donors'], data['late_buy_in_count']), (2, 1, 2))
            self.assertEqual(test.late_buy_in_price, Decimal('320.00'))
            self.assertNotIn('participant_ids', data)

    def test_add_existing_user_uses_same_default_and_override(self):
        self.freeze()
        with self.app.app_context():
            # Add Existing User excludes active pending rows; remove only the fixture pending row.
            db.session.delete(db.session.get(Participation, self.ids['late']))
            db.session.commit()
        self.client.post(f'/admin/add-participant/{self.test_id}', data={'user_id': self.users['late'], 'buy_in_amount': '90'})
        self.assertEqual(self.charges(), {'donor': 280, 'member': 320, 'late': 90})

    def test_reopening_keeps_frozen_until_preview_is_confirmed(self):
        self.freeze()
        self.edit_test('recruiting')
        self.approve('late')
        self.assertEqual(self.charges()['member'], 320)
        page = self.client.get(f'/admin/recalculate-costs/{self.test_id}')
        self.assertIn('Confirm Unlock', page.get_data(as_text=True))
        token = re.search(r'name="preview_token" value="([^"]+)"', page.get_data(as_text=True)).group(1)
        self.client.post(f'/admin/recalculate-costs/{self.test_id}', data={'preview_token': token})
        with self.app.app_context():
            self.assertFalse(db.session.get(GroupTest, self.test_id).costs_locked)
        self.approve('other')
        self.assertEqual(self.charges(), {'donor': 180, 'member': 210, 'late': 320, 'other': 210})
        editor = self.client.get(f'/admin/update-participant/{self.ids["late"]}').get_data(as_text=True)
        self.assertIn('name="buy_in_amount"', editor)
        self.assertIn('Leave blank to keep their assigned charge', editor)

    def test_preview_stale_roster_or_payment_cannot_confirm(self):
        self.freeze()
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            _, token = costs.preview(test)
        self.approve('late')
        response = self.client.post(f'/admin/recalculate-costs/{self.test_id}', data={'preview_token': token}, follow_redirects=True)
        self.assertIn('changed', response.get_data(as_text=True))
        self.assertEqual(self.charges()['member'], 320)
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            _, token = costs.preview(test)
            db.session.get(Participation, self.ids['member']).amount_paid = 10
            db.session.commit()
        response = self.client.post(f'/admin/recalculate-costs/{self.test_id}', data={'preview_token': token}, follow_redirects=True)
        self.assertIn('updated preview', response.get_data(as_text=True))

    def test_legacy_recalc_post_without_preview_does_not_change_charges(self):
        self.freeze()
        self.edit_test('testing', total_lab_cost='1200')
        response = self.client.post(f'/admin/recalculate-costs/{self.test_id}', data={})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.charges()['member'], 320)

    def test_manual_recalculation_outside_recruiting_refreezes(self):
        self.freeze()
        self.approve('late', buy_in_amount='90')
        self.edit_test('testing', total_lab_cost='1200')
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            _, token = costs.preview(test)
        self.client.post(f'/admin/recalculate-costs/{self.test_id}', data={'preview_token': token})
        self.assertEqual(self.charges(), {'donor': 580, 'member': 620, 'late': 90})
        with self.app.app_context():
            self.assertTrue(db.session.get(GroupTest, self.test_id).costs_locked)

    def test_frozen_unapproval_role_change_and_removal_do_not_touch_other_charges(self):
        self.freeze()
        self.approve('late')
        self.client.post(f'/admin/manage-participants/{self.test_id}/unapprove/{self.ids["donor"]}')
        self.assertEqual(self.charges(), {'member': 320, 'late': 320})
        # Reapproving an original funder keeps their original price and cohort membership.
        self.approve('donor')
        self.assertEqual(self.charges()['donor'], 280)
        with self.app.app_context():
            part = db.session.get(Participation, self.ids['donor'])
            self.assertFalse(part.is_late_buy_in)
        self.client.post(f'/admin/remove-participant/{self.ids["donor"]}')
        self.assertEqual(self.charges(), {'member': 320, 'late': 320})

    def test_recruiting_removal_cost_edit_and_donor_change_recalculate(self):
        self.approve('donor')
        self.approve('member')
        self.edit_test('recruiting', total_lab_cost='800')
        self.assertEqual(self.charges(), {'donor': 380, 'member': 420})
        self.client.post(f'/admin/update-participant/{self.ids["donor"]}', data={
            'name': 'donor', 'approved': 'y', 'active': 'y', 'order_status': 'pending', 'amount_paid': '0',
        })
        self.assertEqual(self.charges(), {'donor': 400, 'member': 400})
        self.client.post(f'/admin/remove-participant/{self.ids["donor"]}')
        self.assertEqual(self.charges(), {'member': 800})

    def test_frozen_budget_edits_api_personal_panel_and_balance_use_assigned_price(self):
        self.freeze()
        with self.app.app_context():
            db.session.get(Participation, self.ids['member']).amount_paid = 20
            db.session.commit()
        self.approve('late')
        self.edit_test('testing', total_lab_cost='1200')
        api = self.client.get(f'/api/test/{self.test_id}/costs').json
        self.assertEqual(api['non_donor_pays'], 320)
        self.assertEqual(api['total_fixed_cost'], 600)
        self.login('member')
        page = self.client.get(f'/test/{self.test_id}').get_data(as_text=True)
        self.assertIn('Amount Owed', page)
        self.assertIn('Your amount owed: <strong>$300.00</strong>', page)
        self.assertIn('data-label="Amount Owed"', page)
        self.assertEqual(self.client.get(f'/payments/{self.ids["member"]}').json['amount_owed'], '320.00')

    def test_notifications_bot_status_and_payment_alerts_use_assigned_charge(self):
        self.freeze()
        self.approve('late', buy_in_amount='75')
        with self.app.app_context():
            template = NotificationTemplate(name='Assigned amount', email_body='{{ username }} owes {{ amount_owed }}')
            db.session.add(template)
            db.session.commit()
            template_id = template.id
            test = db.session.get(GroupTest, self.test_id)
            member = db.session.get(User, self.users['member'])
            status = _telegram_status_summary_for_user(member, test)
            self.assertIn('$320.00', status)
        with patch('app.routes.send_group_test_notification') as notify:
            self.client.post(f'/test/{self.test_id}', data={'template_id': template_id})
        amounts = {call.args[1].username: call.kwargs['amount_owed'] for call in notify.call_args_list}
        self.assertEqual(amounts, {'donor': 280, 'member': 320, 'late': 75})
        with self.app.test_request_context(), patch('app.routes.send_notification_message') as notify:
            from app.routes import _notify_payment_claim
            part = db.session.get(Participation, self.ids['member'])
            _notify_payment_claim(part.group_test, part)
            self.assertIn('320.00', notify.call_args.args[3])

    def test_freeze_is_committed_before_status_notification_is_enqueued(self):
        self.approve('donor')
        self.approve('member')
        def observe(test, old):
            self.assertTrue(test.costs_frozen)
            self.assertEqual(test.cost_snapshot['non_donor_pays'], 320)
            self.assertEqual(db.session.get(Participation, self.ids['member']).amount_owed, 320)
        with patch('app.routes._notify_status_change_after_commit', side_effect=observe) as notify:
            self.client.post(f'/admin/set-results/{self.test_id}', data={'results_link': 'https://example.test/report'})
        notify.assert_called_once()

    def test_export_freezes_basis_and_preserves_paid_amounts_and_evidence(self):
        self.freeze()
        with self.app.app_context():
            part = db.session.get(Participation, self.ids['member'])
            part.amount_paid = 320
            part.paid_lab = True
            part.payment_verified_by_id = self.users['admin']
            report = PaymentSubmission(participation_id=part.id, amount=320, method_snapshot={'title': 'Lab account'}, status='approved', transaction_id='original-tx')
            db.session.add(report)
            db.session.commit()
        self.approve('late', buy_in_amount='75')
        self.edit_test('testing', total_lab_cost='1200')
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            workbook = load_workbook(generate_test_export(test), data_only=False)
            sheet = workbook.active
            records = {row[1].value: row for row in sheet.iter_rows(min_row=17, max_row=20) if row[1].value}
            self.assertEqual(records['member'][12].value, 320)
            self.assertEqual(records['member'][13].value, 320)
            self.assertEqual(records['late'][12].value, 75)
            values = {row[0].value: row[1].value for row in sheet.iter_rows() if row[0].value}
            self.assertEqual(values['FROZEN NON-DONOR SHARE'], 320)
            self.assertEqual(values['TOTAL COST-SHARING PARTICIPANTS'], 2)
            part = db.session.get(Participation, self.ids['member'])
            self.assertTrue(part.paid_lab)
            self.assertEqual(part.payment_verified_by_id, self.users['admin'])
            self.assertEqual(part.payment_submissions[0].transaction_id, 'original-tx')

    def test_legacy_ambiguous_charges_are_preserved_until_admin_confirms_buy_in(self):
        with self.app.app_context():
            test = db.session.get(GroupTest, self.test_id)
            test.status = 'closed'
            for name, charge in [('member', 50), ('other', 65)]:
                part = db.session.get(Participation, self.ids[name])
                part.approved = True
                part.amount_owed = charge
            db.session.commit()
            self.assertIsNone(test.late_buy_in_price)
        response = self.approve('late', queue=True)
        self.assertEqual(self.charges(), {'member': 50, 'other': 65})
        self.approve('late', buy_in_amount='60')
        self.assertEqual(self.charges(), {'member': 50, 'other': 65, 'late': 60})
        with self.app.app_context():
            self.assertEqual(db.session.get(GroupTest, self.test_id).late_buy_in_price, Decimal('60'))

    def test_default_change_only_affects_future_buyers_and_invalid_input_rolls_back(self):
        self.freeze()
        for amount in ['NaN', 'Infinity', '-1', '1.234']:
            self.approve('late', buy_in_amount=amount)
            self.assertNotIn('late', self.charges())
        self.client.post(f'/admin/test/{self.test_id}/buy-in-price', data={'buy_in_amount': '95'})
        self.approve('late')
        self.assertEqual(self.charges(), {'donor': 280, 'member': 320, 'late': 95})

    def test_non_admin_cannot_reprice_or_change_default(self):
        self.freeze()
        self.login('member')
        for path in [f'/admin/recalculate-costs/{self.test_id}', f'/admin/test/{self.test_id}/buy-in-price']:
            self.client.post(path, data={'buy_in_amount': '1', 'preview_token': 'anything'})
        self.assertEqual(self.charges(), {'donor': 280, 'member': 320})

    def test_participant_editor_approval_and_frozen_role_edits_use_same_policy(self):
        self.approve('donor')
        payload = {'name': 'member', 'approved': 'y', 'active': 'y',
                   'order_status': 'pending', 'amount_paid': '0'}
        self.client.post(f'/admin/update-participant/{self.ids["member"]}', data=payload)
        self.assertEqual(self.charges(), {'donor': 280, 'member': 320})
        self.edit_test('ready_for_payment')
        payload.update(name='late', buy_in_amount='80', vial_donor='y')
        self.client.post(f'/admin/update-participant/{self.ids["late"]}', data=payload)
        self.assertEqual(self.charges(), {'donor': 280, 'member': 320, 'late': 80})
        payload.update(name='member', buy_in_amount='')
        self.client.post(f'/admin/update-participant/{self.ids["member"]}', data=payload)
        self.assertEqual(self.charges()['member'], 320)

    def test_no_frozen_non_donor_rate_requires_confirmation_and_bulk_is_atomic(self):
        self.approve('donor')
        self.edit_test('ready_for_payment')
        response = self.client.post('/admin/action-queue/approve-selected', data={
            'part_ids': [self.ids['member'], self.ids['late']]}, follow_redirects=True)
        self.assertIn('Set a late buy-in price', response.get_data(as_text=True))
        self.assertEqual(self.charges(), {'donor': 580})
        self.approve('member', buy_in_amount='0')
        self.approve('late')
        self.assertEqual(self.charges(), {'donor': 580, 'member': 0, 'late': 0})

    def test_zero_donor_credit_is_respected_when_recruiting_and_freezing(self):
        self.approve('donor')
        self.approve('member')
        self.edit_test('ready_for_payment', refund_per_donor='0')
        self.assertEqual(self.charges(), {'donor': 300, 'member': 300})

    def test_cost_revision_rejects_concurrent_snapshot_writes(self):
        with self.app.app_context(), Session(db.engine) as one, Session(db.engine) as two:
            first = one.get(GroupTest, self.test_id)
            second = two.get(GroupTest, self.test_id)
            first.cost_revision += 1
            one.commit()
            second.costs_frozen = True
            with self.assertRaises(StaleDataError):
                two.commit()


class CostPolicyMigrationTests(unittest.TestCase):
    def test_upgrade_freezes_only_existing_non_recruiting_without_repricing(self):
        with tempfile.TemporaryDirectory() as directory:
            app = create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': 'sqlite:///' + str(Path(directory) / 'migration.db')})
            with app.app_context():
                try:
                    upgrade(directory='migrations', revision='f8a2c4e6b9d1')
                    db.session.execute(text("INSERT INTO users (id,username,email,password_hash,is_admin,is_active,created_at) VALUES (1,'legacy','legacy@example.test','hash',1,1,CURRENT_TIMESTAMP)"))
                    for ident, status in [(1, 'recruiting'), (2, 'closed')]:
                        db.session.execute(text("INSERT INTO group_tests (id,title,status,total_lab_cost,shipping_cost,donor_shipping_cost,donor_shipping_reimbursement,refund_per_donor,created_by,created_at,updated_at) VALUES (:id,'Legacy',:status,600,0,0,'credit',20,1,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)"), {'id': ident, 'status': status})
                    db.session.execute(text("INSERT INTO participations (group_test_id,user_id,approved,amount_owed,amount_paid,payment_revision,requested_at) VALUES (2,1,1,87.50,60,1,CURRENT_TIMESTAMP)"))
                    db.session.commit()
                    upgrade(directory='migrations')
                    self.assertFalse(db.session.get(GroupTest, 1).costs_locked)
                    frozen = db.session.get(GroupTest, 2)
                    self.assertTrue(frozen.costs_locked)
                    part = Participation.query.one()
                    self.assertEqual((part.amount_owed, part.amount_paid), (87.5, 60))
                    self.assertEqual(frozen.late_buy_in_price, 87.5)
                    self.assertFalse(part.is_late_buy_in)
                    db.session.remove()
                    downgrade(directory='migrations', revision='f8a2c4e6b9d1')
                    self.assertNotIn('costs_frozen', {c['name'] for c in inspect(db.engine).get_columns('group_tests')})
                    self.assertEqual(db.session.execute(text('SELECT amount_owed FROM participations')).scalar_one(), 87.5)
                finally:
                    db.session.remove()
                    db.engine.dispose()
